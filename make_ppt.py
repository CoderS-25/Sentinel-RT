from pptx import Presentation

prs = Presentation()

# Slide 1
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = "Sentinel-RT"
slide.placeholders[1].text = "AI-Driven Real-Time Task Scheduler for μT-Kernel 3.0\nTRON Programming Contest 2026"

# Slide 2
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "The Problem"
slide.placeholders[1].text = "• Static Priority Scheduling is rigid.\n• Modern IoT edge workloads are dynamic and unpredictable.\n• Results in missed deadlines and inefficient power management."

# Slide 3
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Our Solution: TRON x AI"
slide.placeholders[1].text = "• Integrates AI directly into the μT-Kernel 3.0 logic.\n• A 1D-CNN neural network continuously monitors task states.\n• Dynamically re-assigns task priorities on the fly.\n• Intelligently shifts the MCU into power-saving modes."

# Slide 4
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Architecture"
slide.placeholders[1].text = "1. Task Profiler: Gathers CPU load and deadline proximity.\n2. Inference Engine: Runs INT8 TFLite model on the STM32N6 Neural-ART NPU.\n3. Decision Engine: Applies tk_chg_pri() and HAL power states."

# Slide 5
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Performance"
slide.placeholders[1].text = "• Hardware: STM32N6570-DK\n• Model Size: 22.5 KB (INT8 Quantized)\n• Prediction Accuracy: 93.15%\n• Execution: External Flash (XIP)"

# Slide 6
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Thank You"
slide.placeholders[1].text = "Sentinel-RT unlocks the next generation of intelligent RTOS.\nCheck the GitHub repository for the full source code."

prs.save("Sentinel_RT_Presentation.pptx")
